"""States that publish a spreadsheet: one adapter per state, kept small.

Each adapter takes the downloaded bytes and returns Row objects. They live
together here because they are all the same shape - find the header, map a few
columns, emit rows - and because a state that changes its columns should be one
short function to fix, with the state's own file as the only source.
"""
from __future__ import annotations

import csv
import io
from typing import Callable, Iterator

from ..rows import Row, rate

Adapter = Callable[[bytes, str], list[Row]]


def _csv_rows(body: bytes) -> Iterator[dict[str, str]]:
    text = body.decode("utf-8-sig", "replace")
    yield from csv.DictReader(io.StringIO(text))


def _xlsx_rows(body: bytes, sheet: int | str = 0, header_row: int = 1) -> Iterator[dict[str, str]]:
    from openpyxl import load_workbook  # imported here so CSV-only states need no dependency

    wb = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    ws = wb[sheet] if isinstance(sheet, str) else wb[wb.sheetnames[sheet]]
    rows = ws.iter_rows(values_only=True)
    header: list[str] = []
    for i, row in enumerate(rows, start=1):
        values = ["" if v is None else str(v).strip() for v in row]
        if i < header_row:
            continue
        if i == header_row:
            header = [v.strip() for v in values]
            continue
        if not any(values):
            continue
        yield dict(zip(header, values))


def _clean(name: str) -> str:
    return " ".join(str(name).split()).strip()


def texas(body: bytes, source_url: str) -> list[Row]:
    """Comptroller city-rates.xlsx - city code, county, rate."""
    out: list[Row] = []
    for r in _xlsx_rows(body):
        keys = {k.lower(): k for k in r if k}
        code = r.get(keys.get("code", ""), "") or r.get(keys.get("city code", ""), "")
        name = _clean(r.get(keys.get("city", ""), "") or r.get(keys.get("name", ""), ""))
        if not code or not name:
            continue
        out.append(
            Row(
                state="TX",
                jurisdiction_name=name,
                jurisdiction_type="city",
                official_code=str(code).strip(),
                code_system="TX Comptroller local code",
                county_name=_clean(r.get(keys.get("county", ""), "")),
                local_rate=rate(r.get(keys.get("rate", ""), "") or ""),
                source_url=source_url,
            )
        )
    return out


def arizona(body: bytes, source_url: str) -> list[Row]:
    """ADOR TPT rate table - region codes; county rows already include the state rate."""
    out: list[Row] = []
    for r in _csv_rows(body):
        region = (r.get("REGION CODE") or r.get("Region Code") or "").strip()
        name = _clean(r.get("REGION NAME") or r.get("Region Name") or "")
        business_code = (r.get("BUSINESS CODE") or r.get("Business Code") or "").strip()
        if not region or not name or business_code not in ("017", "17", "605"):
            continue  # retail and remote-seller business codes only
        is_county = len(region) == 3 and name.upper().endswith("COUNTY")
        out.append(
            Row(
                state="AZ",
                jurisdiction_name=name.replace(" COUNTY", "").title() if is_county else name.title(),
                jurisdiction_type="county" if is_county else "city",
                official_code=region,
                code_system="AZ ADOR TPT region/city code",
                combined_rate=rate(r.get("TAX RATE") or r.get("Tax Rate") or ""),
                rate_includes_state="true" if is_county else "false",
                source_url=source_url,
                notes=f"business code {business_code}",
            )
        )
    return out


def washington(body: bytes, source_url: str) -> list[Row]:
    """DOR quarterly location-code file - one code per location, not per city."""
    out: list[Row] = []
    for r in _xlsx_rows(body, header_row=4):
        keys = {k.lower(): k for k in r if k}
        code = str(r.get(keys.get("location code", ""), "") or "").strip()
        name = _clean(r.get(keys.get("location", ""), "") or r.get(keys.get("city", ""), ""))
        if not code or not name:
            continue
        kind = "unincorporated" if "unincorp" in name.lower() else "city"
        out.append(
            Row(
                state="WA",
                jurisdiction_name=name,
                jurisdiction_type=kind,
                official_code=code.zfill(4),
                code_system="WA location code",
                county_name=_clean(r.get(keys.get("county", ""), "")),
                local_rate=rate(r.get(keys.get("local rate", ""), "") or ""),
                combined_rate=rate(r.get(keys.get("total rate", ""), "") or ""),
                source_url=source_url,
            )
        )
    return out


ADAPTERS: dict[str, Adapter] = {
    "tx-comptroller": texas,
    "az-adortpt": arizona,
    "wa-dor": washington,
}
